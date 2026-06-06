"""Convert OBJECTIVES_ALIGNMENT.md to a formatted DOCX document."""

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_background(cell, fill_color):
    """Set a table cell background colour (hex string without #)."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_color)
    tcPr.append(shd)


def set_cell_border(cell):
    """Add thin border to a table cell."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        border = OxmlElement(f'w:{side}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), 'CCCCCC')
        tcBorders.append(border)
    tcPr.append(tcBorders)


def add_formatted_run(para, text):
    """Add a run to a paragraph, handling **bold** and *italic* inline markup."""
    parts = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = para.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*') and not part.startswith('**'):
            run = para.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = para.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
        else:
            if part:
                para.add_run(part)


def parse_and_build(md_path: Path, doc: Document):
    lines = md_path.read_text(encoding='utf-8').splitlines()

    # Configure default body style
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    i = 0
    while i < len(lines):
        line = lines[i]

        # ── Heading 1 (# )
        if line.startswith('# ') and not line.startswith('## '):
            para = doc.add_heading(line[2:].strip(), level=1)
            para.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
            i += 1
            continue

        # ── Heading 2 (## )
        if line.startswith('## ') and not line.startswith('### '):
            para = doc.add_heading(line[3:].strip(), level=2)
            para.runs[0].font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
            i += 1
            continue

        # ── Heading 3 (### )
        if line.startswith('### ') and not line.startswith('#### '):
            para = doc.add_heading(line[4:].strip(), level=3)
            para.runs[0].font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
            i += 1
            continue

        # ── Heading 4 (#### )
        if line.startswith('#### '):
            para = doc.add_heading(line[5:].strip(), level=4)
            i += 1
            continue

        # ── Horizontal rule
        if re.match(r'^-{3,}$', line.strip()) or re.match(r'^={3,}$', line.strip()):
            para = doc.add_paragraph()
            pPr = para._p.get_or_add_pPr()
            pBdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:space'), '1')
            bottom.set(qn('w:color'), '2E74B5')
            pBdr.append(bottom)
            pPr.append(pBdr)
            i += 1
            continue

        # ── Table detection (lines starting with |)
        if line.startswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].startswith('|'):
                table_lines.append(lines[i])
                i += 1
            # filter out separator rows (---|---|---)
            data_rows = [r for r in table_lines if not re.match(r'^\|[\s\-:|]+\|', r)]
            if not data_rows:
                continue
            parsed = []
            for row in data_rows:
                cells = [c.strip() for c in row.strip('|').split('|')]
                parsed.append(cells)
            if not parsed:
                continue
            num_cols = max(len(r) for r in parsed)
            tbl = doc.add_table(rows=len(parsed), cols=num_cols)
            tbl.style = 'Table Grid'
            for r_idx, row_data in enumerate(parsed):
                for c_idx in range(num_cols):
                    cell = tbl.cell(r_idx, c_idx)
                    cell_text = row_data[c_idx] if c_idx < len(row_data) else ''
                    # Clear default paragraph and add formatted one
                    cell.paragraphs[0].clear()
                    para = cell.paragraphs[0]
                    if r_idx == 0:
                        set_cell_background(cell, 'DEEAF1')
                        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    add_formatted_run(para, cell_text)
                    if r_idx == 0:
                        for run in para.runs:
                            run.bold = True
                            run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
                    set_cell_border(cell)
            doc.add_paragraph()
            continue

        # ── Bullet list
        if line.startswith('- ') or line.startswith('* '):
            para = doc.add_paragraph(style='List Bullet')
            add_formatted_run(para, line[2:].strip())
            i += 1
            continue

        # ── Numbered list
        m = re.match(r'^\d+\.\s+(.*)', line)
        if m:
            para = doc.add_paragraph(style='List Number')
            add_formatted_run(para, m.group(1).strip())
            i += 1
            continue

        # ── Empty line
        if line.strip() == '':
            i += 1
            continue

        # ── Regular paragraph
        para = doc.add_paragraph()
        add_formatted_run(para, line.strip())
        i += 1

    return doc


def main():
    here = Path(__file__).parent
    md_file = here / 'OBJECTIVES_ALIGNMENT.md'
    out_file = here / 'OBJECTIVES_ALIGNMENT.docx'

    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.5)

    parse_and_build(md_file, doc)
    doc.save(out_file)
    print(f"Saved: {out_file}")


if __name__ == '__main__':
    main()
